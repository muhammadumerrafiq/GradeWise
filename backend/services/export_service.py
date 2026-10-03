"""
GradeWise — Export Service
Generates individual student reports (PDF, DOCX) and batch exports (Excel summary, CSV, Bulk ZIP).
"""

import csv
import io
import os
import re
import zipfile
from datetime import datetime
from typing import Any, Dict, List

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from config import settings


def sanitize_filename(name: str) -> str:
    """Removes special characters to produce safe filename."""
    name = re.sub(r"[^\w\s-]", "", name).strip()
    return re.sub(r"[-\s]+", "_", name)


class ExportService:
    def __init__(self):
        settings.create_directories()

    def get_assignment_export_dir(self, assignment_id: str) -> str:
        d = os.path.join(settings.EXPORT_DIR, str(assignment_id))
        os.makedirs(d, exist_ok=True)
        return d

    def generate_pdf_report(
        self,
        student_name: str,
        assignment_title: str,
        total_score: int,
        max_score: int,
        percentage: float,
        grade_label: str,
        criterion_scores: List[Dict[str, Any]],
        requirements: List[Dict[str, Any]],
        english_analysis: Dict[str, Any],
        strengths: List[str],
        improvements: List[str],
        teacher_feedback: str,
        evaluated_date: str,
        output_path: str,
    ) -> str:
        """Generates a professional PDF evaluation report using ReportLab."""
        doc = SimpleDocTemplate(
            output_path,
            pagesize=letter,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#111827"),
        )
        subtitle_style = ParagraphStyle(
            "DocSubtitle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#6B7280"),
        )
        section_heading = ParagraphStyle(
            "SectionHeading",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#1E3A8A"),
            spaceBefore=10,
            spaceAfter=4,
        )
        body_style = ParagraphStyle(
            "Body",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=14,
            textColor=colors.HexColor("#1F2937"),
        )
        bullet_style = ParagraphStyle(
            "Bullet",
            parent=body_style,
            leftIndent=12,
            firstLineIndent=-12,
            spaceAfter=3,
        )

        story = []

        # Header Table
        header_data = [
            [
                Paragraph("<b>Assignment Evaluation Report</b>", title_style),
                Paragraph(f"<b>Date:</b> {evaluated_date}", subtitle_style),
            ],
            [
                Paragraph(f"Assignment: <b>{assignment_title}</b>", subtitle_style),
                Paragraph(f"GradeWise Assessment", subtitle_style),
            ]
        ]
        header_table = Table(header_data, colWidths=[380, 150])
        header_table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ]))
        story.append(header_table)
        story.append(Spacer(1, 10))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#2563EB"), spaceAfter=10))

        # Student & Score Summary Box
        grade_display = grade_label if grade_label else "N/A"
        summary_data = [
            [
                Paragraph(f"<b>Student:</b> {student_name}", body_style),
                Paragraph(f"<b>Score:</b> {total_score}/{max_score} ({percentage:.1f}%)", body_style),
                Paragraph(f"<b>Grade:</b> {grade_display}", body_style),
            ]
        ]
        summary_table = Table(summary_data, colWidths=[200, 200, 130])
        summary_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EFF6FF")),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#BFDBFE")),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ]))
        story.append(summary_table)
        story.append(Spacer(1, 12))

        # Rubric Breakdown
        story.append(Paragraph("RUBRIC BREAKDOWN", section_heading))
        rubric_rows = [["Criterion", "Score", "Performance", "Rationale / Evidence"]]
        for c in criterion_scores:
            name = c.get("criterion_name", "")
            s = c.get("score", 0)
            m = c.get("max_score", 1)
            pct = (s / m) if m > 0 else 0
            filled = int(round(pct * 10))
            bar_text = "[" + "■" * filled + "□" * (10 - filled) + "]"
            rationale_text = c.get("rationale", "")
            evidence_val = c.get("evidence", "")
            quote_text = f'<i>Quote: "{evidence_val}"</i>' if evidence_val else ""
            detail = f"{rationale_text} {quote_text}".strip()
            rubric_rows.append([
                Paragraph(f"<b>{name}</b>", body_style),
                Paragraph(f"{s}/{m}", body_style),
                Paragraph(f"<font color='#2563EB'>{bar_text}</font>", body_style),
                Paragraph(detail, body_style),
            ])

        rubric_table = Table(rubric_rows, colWidths=[120, 50, 90, 270])
        rubric_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F4F6")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#374151")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(rubric_table)
        story.append(Spacer(1, 10))

        # Requirements
        if requirements:
            story.append(Paragraph("CHECKLIST & REQUIREMENTS", section_heading))
            req_rows = []
            for r in requirements:
                st = r.get("status", "PARTIAL").upper()
                desc = r.get("description", "")
                detail = r.get("detail", "")
                if st == "PASS":
                    icon = "<font color='#16A34A'><b>✓ PASS</b></font>"
                elif st == "FAIL":
                    icon = "<font color='#DC2626'><b>✗ FAIL</b></font>"
                else:
                    icon = "<font color='#D97706'><b>◐ PARTIAL</b></font>"
                req_rows.append([
                    Paragraph(icon, body_style),
                    Paragraph(f"<b>{desc}</b> — {detail}", body_style)
                ])
            req_table = Table(req_rows, colWidths=[80, 450])
            req_table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(req_table)
            story.append(Spacer(1, 10))

        # English Analysis
        if english_analysis:
            story.append(Paragraph("ENGLISH LANGUAGE ANALYSIS", section_heading))
            analysis_rows = [["Category", "Severity", "Analysis & Issues"]]
            for cat in ["grammar", "vocabulary", "tenses", "mechanics", "structure"]:
                data = english_analysis.get(cat, {})
                summ = data.get("summary", "No issues identified.")
                issues = data.get("issues", [])
                issues_text = f"<br/>• " + "<br/>• ".join(issues) if issues else ""
                sev = data.get("severity", "none").upper()
                color_map = {
                    "NONE": "#4B5563",
                    "MINOR": "#2563EB",
                    "MODERATE": "#D97706",
                    "SIGNIFICANT": "#DC2626",
                }
                sev_color = color_map.get(sev, "#4B5563")
                analysis_rows.append([
                    Paragraph(f"<b>{cat.capitalize()}</b>", body_style),
                    Paragraph(f"<font color='{sev_color}'><b>{sev}</b></font>", body_style),
                    Paragraph(f"{summ}{issues_text}", body_style),
                ])
            eng_table = Table(analysis_rows, colWidths=[100, 80, 350])
            eng_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F4F6")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(eng_table)
            story.append(Spacer(1, 10))

        # Strengths & Improvements
        if strengths or improvements:
            feedback_table_data = []
            col_left = [Paragraph("<b>STRENGTHS</b>", section_heading)]
            for s in strengths:
                col_left.append(Paragraph(f"• {s}", bullet_style))

            col_right = [Paragraph("<b>AREAS FOR IMPROVEMENT</b>", section_heading)]
            for imp in improvements:
                col_right.append(Paragraph(f"• {imp}", bullet_style))

            si_table = Table([[col_left, col_right]], colWidths=[265, 265])
            si_table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ]))
            story.append(si_table)
            story.append(Spacer(1, 10))

        # Teacher Feedback
        story.append(Paragraph("TEACHER FEEDBACK", section_heading))
        feedback_box = Table(
            [[Paragraph(teacher_feedback.replace("\n", "<br/>"), body_style)]],
            colWidths=[530],
        )
        feedback_box.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FAFAFA")),
            ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#E5E7EB")),
            ("TOPPADDING", (0, 0), (-1, -1), 8),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ("LEFTPADDING", (0, 0), (-1, -1), 10),
            ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ]))
        story.append(feedback_box)
        story.append(Spacer(1, 15))

        # Footer Date
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#D1D5DB"), spaceAfter=6))
        story.append(Paragraph(f"<i>Evaluated on: {evaluated_date} — GradeWise English Evaluator</i>", subtitle_style))

        doc.build(story)
        return output_path

    def generate_docx_report(
        self,
        student_name: str,
        assignment_title: str,
        total_score: int,
        max_score: int,
        percentage: float,
        grade_label: str,
        criterion_scores: List[Dict[str, Any]],
        requirements: List[Dict[str, Any]],
        english_analysis: Dict[str, Any],
        strengths: List[str],
        improvements: List[str],
        teacher_feedback: str,
        evaluated_date: str,
        output_path: str,
    ) -> str:
        """Generates an individual student evaluation report in Word DOCX format."""
        doc = docx.Document()

        # Document Header
        h = doc.add_heading("Assignment Evaluation Report", level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        doc.add_paragraph(f"Assignment: {assignment_title} | Date: {evaluated_date}")

        # Summary Table
        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = f"Student: {student_name}"
        hdr_cells[1].text = f"Score: {total_score}/{max_score} ({percentage:.1f}%)"
        hdr_cells[2].text = f"Grade: {grade_label or 'N/A'}"

        doc.add_paragraph()

        # Rubric Breakdown
        doc.add_heading("Rubric Breakdown", level=2)
        r_table = doc.add_table(rows=1, cols=3)
        r_hdr = r_table.rows[0].cells
        r_hdr[0].text = "Criterion"
        r_hdr[1].text = "Score"
        r_hdr[2].text = "Rationale & Evidence"
        for c in criterion_scores:
            row_cells = r_table.add_row().cells
            row_cells[0].text = c.get("criterion_name", "")
            row_cells[1].text = f"{c.get('score', 0)} / {c.get('max_score', 0)}"
            rationale = c.get("rationale", "")
            evidence = f" (Quote: \"{c.get('evidence')}\")" if c.get("evidence") else ""
            row_cells[2].text = f"{rationale}{evidence}"

        # Requirements
        if requirements:
            doc.add_heading("Requirements Checklist", level=2)
            for r in requirements:
                st = r.get("status", "PARTIAL")
                desc = r.get("description", "")
                detail = r.get("detail", "")
                doc.add_paragraph(f"[{st}] {desc} — {detail}", style="List Bullet")

        # English Analysis
        if english_analysis:
            doc.add_heading("English Language Analysis", level=2)
            for cat, data in english_analysis.items():
                p = doc.add_paragraph()
                p.add_run(f"{cat.capitalize()} ({data.get('severity', 'none').upper()}): ").bold = True
                p.add_run(data.get("summary", ""))
                for issue in data.get("issues", []):
                    doc.add_paragraph(issue, style="List Bullet 2")

        # Strengths
        if strengths:
            doc.add_heading("Strengths", level=2)
            for s in strengths:
                doc.add_paragraph(s, style="List Bullet")

        # Areas for Improvement
        if improvements:
            doc.add_heading("Areas for Improvement", level=2)
            for imp in improvements:
                doc.add_paragraph(imp, style="List Bullet")

        # Teacher Feedback
        doc.add_heading("Teacher Feedback", level=2)
        doc.add_paragraph(teacher_feedback)

        doc.save(output_path)
        return output_path

    def generate_excel_summary(
        self,
        assignment_title: str,
        rubric_criteria_names: List[str],
        results: List[Dict[str, Any]],
        output_path: str,
    ) -> str:
        """
        Generates Excel summary spreadsheet:
        Student Name | Score | Max Score | Percentage | Grade | [Criteria cols] | Status | Feedback
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Evaluation Summary"

        # Headers
        headers = ["Student Name", "Score", "Max Score", "Percentage", "Grade"]
        for cn in rubric_criteria_names:
            headers.append(cn)
        headers.extend(["Status", "Teacher Feedback"])

        ws.append(headers)

        # Style header row
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data Rows
        for r in results:
            row_vals = [
                r.get("student_name", "Unknown"),
                r.get("total_score", 0),
                r.get("max_score", 0),
                f"{r.get('percentage', 0):.1f}%",
                r.get("grade_label", ""),
            ]
            scores_by_name = {c["criterion_name"]: f"{c['score']}/{c['max_score']}" for c in r.get("criterion_scores", [])}
            for cn in rubric_criteria_names:
                row_vals.append(scores_by_name.get(cn, "N/A"))

            fb = r.get("teacher_feedback", "")
            truncated_fb = (fb[:200] + "...") if len(fb) > 200 else fb

            row_vals.append(r.get("eval_status", "pending_review"))
            row_vals.append(truncated_fb)

            ws.append(row_vals)

        # Adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

        wb.save(output_path)
        return output_path

    def generate_csv_summary(
        self,
        rubric_criteria_names: List[str],
        results: List[Dict[str, Any]],
        output_path: str,
    ) -> str:
        """Generates CSV summary export."""
        headers = ["Student Name", "Score", "Max Score", "Percentage", "Grade"]
        for cn in rubric_criteria_names:
            headers.append(f"{cn} (Score/Max)")
        headers.extend(["Status", "Teacher Feedback"])

        with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for r in results:
                row_vals = [
                    r.get("student_name", "Unknown"),
                    r.get("total_score", 0),
                    r.get("max_score", 0),
                    f"{r.get('percentage', 0):.1f}%",
                    r.get("grade_label", ""),
                ]
                scores_by_name = {c["criterion_name"]: f"{c['score']}/{c['max_score']}" for c in r.get("criterion_scores", [])}
                for cn in rubric_criteria_names:
                    row_vals.append(scores_by_name.get(cn, "N/A"))

                fb = r.get("teacher_feedback", "")
                truncated_fb = (fb[:200] + "...") if len(fb) > 200 else fb

                row_vals.append(r.get("eval_status", "pending_review"))
                row_vals.append(truncated_fb)
                writer.writerow(row_vals)

        return output_path

    def create_bulk_zip(
        self,
        files_to_zip: List[str],
        zip_output_path: str,
    ) -> str:
        """Packages a list of file paths into a single zip archive."""
        with zipfile.ZipFile(zip_output_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for fpath in files_to_zip:
                if os.path.exists(fpath):
                    arcname = os.path.basename(fpath)
                    zf.write(fpath, arcname=arcname)
        return zip_output_path

    def generate_manual_session_csv(
        self,
        session_title: str,
        evaluations: List[Any],
        output_path: str,
    ) -> str:
        """Generates a CSV export of manual paste evaluations."""
        headers = [
            "Student Name",
            "Evaluated Date",
            "Word Count",
            "Requirements Result",
            "Writing Quality Summary",
            "Strengths",
            "Areas for Improvement",
            "Teacher Feedback",
        ]
        with open(output_path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for ev in evaluations:
                name = getattr(ev, "student_name", "") or ""
                date_str = ev.created_at.strftime("%Y-%m-%d %H:%M") if hasattr(ev, "created_at") and ev.created_at else ""
                wc = getattr(ev, "word_count", 0) or 0
                reqs = getattr(ev, "requirements_result", []) or []
                req_summary = "; ".join([f"{r.get('description', '')}: {r.get('status', '')}" for r in reqs if isinstance(r, dict)])
                wq = getattr(ev, "writing_quality_summary", "") or ""
                strengths = "; ".join(getattr(ev, "strengths", []) or [])
                improvements = "; ".join(getattr(ev, "improvements", []) or [])
                fb = getattr(ev, "teacher_feedback", "") or ""
                writer.writerow([name, date_str, wc, req_summary, wq, strengths, improvements, fb])
        return output_path

    def generate_manual_session_docx(
        self,
        session_title: str,
        assignment_type: str,
        evaluations: List[Any],
        output_path: str,
    ) -> str:
        """Generates a formatted Word document for manual paste session."""
        doc = docx.Document()
        title_p = doc.add_paragraph()
        run = title_p.add_run(f"Manual Evaluation Report — {session_title}")
        run.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

        sub_p = doc.add_paragraph()
        sub_p.add_run(f"Type: {assignment_type.title()}  |  Total Evaluated: {len(evaluations)}  |  Generated: {datetime.now().strftime('%B %d, %Y')}")
        sub_p.runs[0].font.size = Pt(10)
        sub_p.runs[0].font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        for ev in evaluations:
            doc.add_heading(f"Student: {getattr(ev, 'student_name', 'Student')}", level=2)
            wc = getattr(ev, "word_count", 0) or 0
            doc.add_paragraph(f"Word count: {wc} words")

            reqs = getattr(ev, "requirements_result", []) or []
            if reqs:
                doc.add_heading("Requirements Check", level=3)
                for r in reqs:
                    if isinstance(r, dict):
                        doc.add_paragraph(f"[{r.get('status', 'PASS')}] {r.get('description', '')} — {r.get('detail', '')}", style="List Bullet")

            wq = getattr(ev, "writing_quality_summary", "")
            if wq:
                doc.add_heading("Writing Quality Summary", level=3)
                doc.add_paragraph(wq)

            strengths = getattr(ev, "strengths", []) or []
            if strengths:
                doc.add_heading("Strengths", level=3)
                for s in strengths:
                    doc.add_paragraph(str(s), style="List Bullet")

            improvements = getattr(ev, "improvements", []) or []
            if improvements:
                doc.add_heading("Areas for Improvement", level=3)
                for imp in improvements:
                    doc.add_paragraph(str(imp), style="List Bullet")

            fb = getattr(ev, "teacher_feedback", "") or ""
            if fb:
                doc.add_heading("Teacher Feedback", level=3)
                fb_p = doc.add_paragraph(fb)
                fb_p.paragraph_format.left_indent = Inches(0.25)

            doc.add_paragraph("—" * 30)

        doc.save(output_path)
        return output_path

    def generate_manual_session_pdf(
        self,
        session_title: str,
        assignment_type: str,
        evaluations: List[Any],
        output_path: str,
    ) -> str:
        """Generates a clean PDF document for manual paste session."""
        from reportlab.platypus import PageBreak

        doc = SimpleDocTemplate(
            output_path,
            pagesize=letter,
            rightMargin=40,
            leftMargin=40,
            topMargin=40,
            bottomMargin=40,
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "MTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=16,
            leading=20,
            textColor=colors.HexColor("#111827"),
        )
        subtitle_style = ParagraphStyle(
            "MSubtitle",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#6B7280"),
        )
        h2_style = ParagraphStyle(
            "MH2",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#1E3A8A"),
            spaceBefore=8,
            spaceAfter=3,
        )
        body_style = ParagraphStyle(
            "MBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor("#1F2937"),
        )
        fb_style = ParagraphStyle(
            "MFB",
            parent=body_style,
            fontName="Helvetica",
            fontSize=9.5,
            leading=14,
            textColor=colors.HexColor("#111827"),
        )

        story = []
        for idx, ev in enumerate(evaluations):
            if idx > 0:
                story.append(PageBreak())

            sname = getattr(ev, "student_name", "Student")
            wc = getattr(ev, "word_count", 0) or 0
            header_table = Table([
                [
                    Paragraph(f"<b>{session_title}</b>", title_style),
                    Paragraph(f"<b>Student:</b> {sname}", subtitle_style),
                ],
                [
                    Paragraph(f"Type: {assignment_type.title()} | {wc} words", subtitle_style),
                    Paragraph(f"Date: {datetime.now().strftime('%Y-%m-%d')}", subtitle_style),
                ]
            ], colWidths=[330, 200])
            header_table.setStyle(TableStyle([
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (1, 0), (1, -1), "RIGHT"),
            ]))
            story.append(header_table)
            story.append(Spacer(1, 6))
            story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#2563EB"), spaceAfter=8))

            reqs = getattr(ev, "requirements_result", []) or []
            if reqs:
                story.append(Paragraph("<b>Requirements Check</b>", h2_style))
                req_rows = [["Requirement", "Status", "Detail"]]
                for r in reqs:
                    if isinstance(r, dict):
                        req_rows.append([
                            Paragraph(r.get("description", ""), body_style),
                            r.get("status", "PASS"),
                            Paragraph(r.get("detail", ""), body_style),
                        ])
                rt = Table(req_rows, colWidths=[200, 60, 270])
                rt.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F3F4F6")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#E5E7EB")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]))
                story.append(rt)
                story.append(Spacer(1, 6))

            wq = getattr(ev, "writing_quality_summary", "")
            if wq:
                story.append(Paragraph("<b>Writing Quality Summary</b>", h2_style))
                story.append(Paragraph(wq, body_style))
                story.append(Spacer(1, 4))

            strengths = getattr(ev, "strengths", []) or []
            improvements = getattr(ev, "improvements", []) or []
            if strengths or improvements:
                story.append(Paragraph("<b>Observations</b>", h2_style))
                for s in strengths:
                    story.append(Paragraph(f"• <b>Strength:</b> {s}", body_style))
                for imp in improvements:
                    story.append(Paragraph(f"• <b>Improvement:</b> {imp}", body_style))
                story.append(Spacer(1, 4))

            fb = getattr(ev, "teacher_feedback", "") or ""
            if fb:
                story.append(Paragraph("<b>Teacher Feedback</b>", h2_style))
                fb_box = Table([[Paragraph(fb, fb_style)]], colWidths=[530])
                fb_box.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F9FAFB")),
                    ("BOX", (0, 0), (-1, -1), 1, colors.HexColor("#D1D5DB")),
                    ("PADDING", (0, 0), (-1, -1), 8),
                ]))
                story.append(fb_box)

        doc.build(story)
        return output_path


export_service = ExportService()


def generate_pdf(evaluation: Any, student_name: str, assignment: Any, output_path: Optional[str] = None) -> str:
    """Convenience module function to generate a PDF evaluation report."""
    if not output_path:
        out_dir = export_service.get_assignment_export_dir(getattr(assignment, "id", "default"))
        output_path = os.path.join(out_dir, f"{sanitize_filename(student_name)}_report.pdf")

    a_title = getattr(assignment, "title", "Assignment")
    tot = getattr(evaluation, "total_score", 0)
    mx = getattr(evaluation, "max_score", getattr(assignment, "total_marks", 100))
    pct = getattr(evaluation, "percentage", 0.0)
    grd = getattr(evaluation, "grade_label", "")
    crit_scores = getattr(evaluation, "criterion_scores", []) or []
    reqs = getattr(evaluation, "requirements_result", []) or []
    eng = getattr(evaluation, "english_analysis", {}) or {}
    strng = getattr(evaluation, "strengths", []) or []
    imp = getattr(evaluation, "improvements", []) or []
    tfb = getattr(evaluation, "teacher_feedback", "") or ""
    dt = datetime.now().strftime("%B %d, %Y")

    return export_service.generate_pdf_report(
        student_name=student_name,
        assignment_title=a_title,
        total_score=tot,
        max_score=mx,
        percentage=pct,
        grade_label=grd,
        criterion_scores=crit_scores if isinstance(crit_scores, list) else [],
        requirements=reqs if isinstance(reqs, list) else [],
        english_analysis=eng if isinstance(eng, dict) else {},
        strengths=strng if isinstance(strng, list) else [],
        improvements=imp if isinstance(imp, list) else [],
        teacher_feedback=tfb,
        evaluated_date=dt,
        output_path=output_path,
    )


def generate_docx(evaluation: Any, student_name: str, assignment: Any, output_path: Optional[str] = None) -> str:
    """Convenience module function to generate a DOCX evaluation report."""
    if not output_path:
        out_dir = export_service.get_assignment_export_dir(getattr(assignment, "id", "default"))
        output_path = os.path.join(out_dir, f"{sanitize_filename(student_name)}_report.docx")

    a_title = getattr(assignment, "title", "Assignment")
    tot = getattr(evaluation, "total_score", 0)
    mx = getattr(evaluation, "max_score", getattr(assignment, "total_marks", 100))
    pct = getattr(evaluation, "percentage", 0.0)
    grd = getattr(evaluation, "grade_label", "")
    crit_scores = getattr(evaluation, "criterion_scores", []) or []
    reqs = getattr(evaluation, "requirements_result", []) or []
    eng = getattr(evaluation, "english_analysis", {}) or {}
    strng = getattr(evaluation, "strengths", []) or []
    imp = getattr(evaluation, "improvements", []) or []
    tfb = getattr(evaluation, "teacher_feedback", "") or ""
    dt = datetime.now().strftime("%B %d, %Y")

    return export_service.generate_docx_report(
        student_name=student_name,
        assignment_title=a_title,
        total_score=tot,
        max_score=mx,
        percentage=pct,
        grade_label=grd,
        criterion_scores=crit_scores if isinstance(crit_scores, list) else [],
        requirements=reqs if isinstance(reqs, list) else [],
        english_analysis=eng if isinstance(eng, dict) else {},
        strengths=strng if isinstance(strng, list) else [],
        improvements=imp if isinstance(imp, list) else [],
        teacher_feedback=tfb,
        evaluated_date=dt,
        output_path=output_path,
    )


def generate_excel_summary(assignment_title: str, rubric_criteria_names: List[str], results: List[Dict[str, Any]], output_path: str) -> str:
    return export_service.generate_excel_summary(assignment_title, rubric_criteria_names, results, output_path)


def generate_csv_summary(rubric_criteria_names: List[str], results: List[Dict[str, Any]], output_path: str) -> str:
    return export_service.generate_csv_summary(rubric_criteria_names, results, output_path)


def generate_bulk_zip(file_paths: List[str], output_path: str) -> str:
    return export_service.generate_bulk_zip(file_paths, output_path)

