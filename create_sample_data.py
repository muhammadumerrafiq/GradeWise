import os
import zipfile
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

SAMPLE_DIR = os.path.join(os.path.dirname(__file__), "sample_submissions")
os.makedirs(SAMPLE_DIR, exist_ok=True)

essays = [
    (
        "Sara_Ahmed_Narrative_Essay.txt",
        """The sound of rain tapping against the library window was the only rhythm that kept me grounded that evening. As the clock struck eight, I realized I had only two hours left before the submission portal locked. My hands trembled as I typed out the final reflections on my journey through high school. 

Every mistake, every triumph, and every late night spent pondering philosophical questions had led to this very crossroad. I looked across the quiet room and saw my grandfather's advice etched in memory: "Never fear the blank page, for it is merely waiting for your courage." Taking a deep breath, I organized the central thesis of my life's story, weaving together themes of perseverance, curiosity, and empathy. The conclusion emerged not as an end, but as a gateway to what lay beyond."""
    ),
    (
        "Marcus_Johnson_English_Draft.txt",
        """In the heart of the metropolis, silence is a luxury seldom afforded. I walked through the neon-lit alleys of Chicago, listening to the ambient symphony of sirens, distant subway rumblings, and the quiet murmurs of weary commuters. This urban ecosystem, vibrant yet unforgiving, taught me more about human resilience than any textbook ever could.

Growing up between two contrasting worlds—the structured routine of academics and the unpredictable dynamism of the city streets—forged my analytical perspective. I observed how language adapts, how slang breathes life into formal discourse, and how stories bridge generational divides. Ultimately, storytelling is our most potent tool for civic empathy."""
    ),
    (
        "unknown_task_submission_991823.txt",
        """Technology has altered the landscape of modern education in profound and irrevocable ways. While critics lament the loss of traditional handwriting and attention spans, digital literacy has democratized access to global information repositories. In this essay, I examine three critical paradigms: remote asynchronous learning, algorithmic grading systems, and personalized adaptive pedagogy.

Evidence suggests that when utilized responsibly, instructional technology fosters self-regulated learning. However, equitable hardware access remains an urgent ethical challenge that educators and policymakers must address collectively."""
    ),
]

# Write TXT files
for filename, text in essays:
    path = os.path.join(SAMPLE_DIR, filename)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text.strip())
    print(f"Created {path}")

# Write DOCX file
docx_name = "David_Chen_Narrative_Submission.docx"
docx_path = os.path.join(SAMPLE_DIR, docx_name)
doc = docx.Document()
doc.add_heading("Navigating Uncertainty: A Personal Reflection", 0)
doc.add_paragraph(
    "Standing at the edge of the harbor, I watched the morning fog slowly dissipate over the cold Atlantic waters. "
    "Uncertainty had always been my greatest adversary, yet this journey required me to embrace the unknown. "
    "Throughout my academic pursuits, I sought predictability in formulas and historical dates, but real life rarely adheres to rigid outlines."
)
doc.add_paragraph(
    "By challenging my innate desire for control, I discovered that adaptability is the cornerstone of intellectual growth. "
    "When our science expedition faced unexpected equipment failures in northern Maine, it was creative problem-solving—not memorized protocols—that saved our project. "
    "This experience reshaped my philosophy on collaboration and persistence."
)
doc.save(docx_path)
print(f"Created {docx_path}")

# Write PDF file
pdf_name = "Emily_Watson_FA23_014_Essay.pdf"
pdf_path = os.path.join(SAMPLE_DIR, pdf_name)
pdf_doc = SimpleDocTemplate(pdf_path, pagesize=letter)
styles = getSampleStyleSheet()
story = [
    Paragraph("<b>The Architecture of Memory</b>", styles["Title"]),
    Spacer(1, 12),
    Paragraph("By Emily Watson (ID: FA23-014)", styles["Normal"]),
    Spacer(1, 14),
    Paragraph(
        "Memory is not a static recording device; rather, it is an active architectural endeavor that continuously reconstructs our past. "
        "In literature, Virginia Woolf and Marcel Proust demonstrated how sensory triggers unlock involuntary recollections. "
        "Through an examination of modern cognitive science alongside modernist prose, we discover that our identities are woven from shifting narrative threads.",
        styles["Normal"]
    ),
    Spacer(1, 10),
    Paragraph(
        "Furthermore, collective memory shapes societal values and historical accountability. "
        "When communities preserve oral histories, they safeguard cultural resilience against systemic erasure. "
        "In conclusion, recognizing the fluidity of memory allows us to approach both literature and human experience with humility and critical awareness.",
        styles["Normal"]
    ),
]
pdf_doc.build(story)
print(f"Created {pdf_path}")

# Create ZIP archive containing all of the above
zip_path = os.path.join(SAMPLE_DIR, "all_student_submissions.zip")
with zipfile.ZipFile(zip_path, "w") as zf:
    for f in os.listdir(SAMPLE_DIR):
        if f.endswith((".txt", ".docx", ".pdf")):
            full_p = os.path.join(SAMPLE_DIR, f)
            zf.write(full_p, arcname=f)
print(f"Created sample ZIP archive: {zip_path}")
